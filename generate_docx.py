from docx import Document
from docx.shared import Pt
import os

with open("Data_Science_Documentation.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()
doc.add_heading("Diet Impact Analyser: Data Science Implementation Documentation", 0)

in_code_block = False
code_text = ""

for line in lines:
    line = line.strip('\n')
    
    if line.startswith("```"):
        if in_code_block:
            # End code block
            p = doc.add_paragraph(code_text)
            p.style = 'Intense Quote'
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
            in_code_block = False
            code_text = ""
        else:
            in_code_block = True
    elif in_code_block:
        code_text += line + "\n"
    else:
        if line.startswith("# "):
            pass # Skip title as it's added up top
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("---"):
            pass
        elif line == "":
            pass
        else:
            p = doc.add_paragraph()
            # Basic bold parsing for **text**
            parts = line.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0:
                    run.bold = True

doc.save("Data_Science_Concepts.docx")
print("Successfully generated Data_Science_Concepts.docx")
